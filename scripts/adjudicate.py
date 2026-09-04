#!/usr/bin/env python3
"""Build a hand-adjudication worksheet: verify ground truth against the source images.

    python3 scripts/adjudicate.py --doc-type invoice --dataset docile100 --n 40

Why this exists. DocILE100's labels come from a Hugging Face upload whose dataset card is
empty — no annotation methodology, no licence, no attribution. The labels look good (69 of 100
documents have line-item totals summing exactly to a stated subtotal or total, and three images
checked by hand were read faithfully), but "looks good" is not a basis for a published number.
This sheet turns a sample of them into verified ground truth.

Output: one xlsx per dataset with three sheets.

  README        what to do, and what the verdicts mean
  fields        one row per (document, scalar field). Confirm, correct, or mark not-on-page.
  line_items    one row per (document, GT line-item row). Same three verdicts, plus a column
                for rows that are ON the page but MISSING from the ground truth — the failure
                the FATURA line-item audit found and the one this dataset is most exposed to.

The sample is stratified by table size, not random: a random 40 of this corpus would be mostly
single-row tables and would say nothing about the 44-row broadcast log, which is exactly where
a line-item extractor breaks.
"""
from __future__ import annotations

import argparse
import collections
import json
import pathlib
import random
import sys

_ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_ROOT))

import doctypes                                                    # noqa: E402
from core.canonical import read_jsonl                              # noqa: E402
from core.normalize import _Absent                                 # noqa: E402
from registry import dataset_for, gt_dir as _gt_dir                # noqa: E402

VERDICTS = ["", "ok", "corrected", "not-on-page", "unsure"]
HELP = {
    "ok": "the ground-truth value is exactly what the page prints",
    "corrected": "the page prints something else — put it in the 'correct value' column",
    "not-on-page": "the page does not print this at all; the label should not exist",
    "unsure": "cannot tell from the image; excluded from the verified subset",
}


def _band(n: int) -> str:
    return ("none" if n == 0 else "single" if n == 1 else "small" if n <= 3
            else "medium" if n <= 9 else "large")


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--doc-type", default="invoice")
    ap.add_argument("--dataset", default=None)
    ap.add_argument("--n", type=int, default=40, help="documents to adjudicate")
    ap.add_argument("--seed", type=int, default=20260903)
    ap.add_argument("-o", "--out", default=None)
    a = ap.parse_args(argv)

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    spec = doctypes.get(a.doc_type)
    entry = dataset_for(spec.name, a.dataset)
    gt_path = _gt_dir(spec.name, entry.name) / "ground_truth.jsonl"
    if not gt_path.exists():
        print(f"!! no ground truth at {gt_path}. Run scripts/build_gt.py first.")
        return 2
    records = {r.doc_id: r for r in read_jsonl(str(gt_path))}
    adapter = entry.load()

    # ---- stratified sample --------------------------------------------------
    bands = collections.defaultdict(list)
    for r in records.values():
        bands[_band(int(r.meta.get("n_gt_rows") or 0))].append(r.doc_id)
    rng = random.Random(a.seed)
    # Every band gets weight; the big-table bands get all of what little they have, because
    # they are the cases that decide whether a line-item number means anything.
    quota = {"large": 10, "medium": 10, "small": 8, "single": 8, "none": 4}
    chosen: list = []
    for band in ("large", "medium", "small", "single", "none"):
        ids = sorted(bands.get(band, []))
        take = min(quota[band], len(ids))
        chosen += rng.sample(ids, take)
    # top up from the largest tables if a band was short
    if len(chosen) < a.n:
        rest = sorted(set(records) - set(chosen),
                      key=lambda d: -int(records[d].meta.get("n_gt_rows") or 0))
        chosen += rest[: a.n - len(chosen)]
    chosen = sorted(chosen[: a.n], key=lambda d: -int(records[d].meta.get("n_gt_rows") or 0))

    out = pathlib.Path(a.out) if a.out else (_gt_dir(spec.name, entry.name)
                                             / f"adjudication_{entry.name}_{len(chosen)}.xlsx")
    HDR = PatternFill("solid", fgColor="1F3864")
    HDRF = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    BODY = Font(name="Arial", size=10)
    MONO = Font(name="Consolas", size=9)
    INPUT = PatternFill("solid", fgColor="FFFF00")

    wb = Workbook()

    # ---- README -------------------------------------------------------------
    ws = wb.active
    ws.title = "README"
    lines = [
        ("How to use this sheet", True),
        ("", False),
        (f"{len(chosen)} of {len(records)} {entry.name} documents, stratified by line-item "
         f"table size so the large tables are over-represented on purpose.", False),
        ("", False),
        ("Open each document's image and check the ground truth against what the page prints. "
         "Fill the YELLOW columns only.", False),
        ("", False),
        ("verdict values:", True),
    ]
    for k, v in HELP.items():
        lines.append((f"    {k:14s} {v}", False))
    lines += [
        ("", False),
        ("On the line_items sheet there is one extra thing to record: if the page shows rows "
         "that are NOT in this sheet, put how many in 'rows missing from GT' on the FIRST row "
         "of that document. A ground truth that silently under-reports rows makes a model's "
         "row recall look better than it is — that is the exact defect found in FATURA's line "
         "items, and it is the one this dataset is most exposed to.", False),
        ("", False),
        ("A document counts as VERIFIED only when every one of its rows on both sheets is "
         "marked ok or corrected. Anything left blank or marked unsure keeps the document out "
         "of the published subset.", False),
        ("", False),
        (f"images: {adapter.root / 'images'}", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = Font(name="Arial", size=11, bold=bold)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120

    def header(sheet, cols):
        sheet.append(cols)
        for i in range(1, len(cols) + 1):
            c = sheet.cell(row=1, column=i)
            c.fill, c.font = HDR, HDRF
        sheet.freeze_panes = "A2"

    def finish(sheet, widths, yellow, n_rows):
        for i, w in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(i)].width = w
        dv = DataValidation(type="list", formula1='"' + ",".join(VERDICTS[1:]) + '"',
                            allow_blank=True)
        sheet.add_data_validation(dv)
        for col in yellow:
            L = get_column_letter(col)
            for r in range(2, n_rows + 2):
                sheet.cell(row=r, column=col).fill = INPUT
            if col == yellow[0]:
                dv.add(f"{L}2:{L}{n_rows + 1}")

    # ---- fields -------------------------------------------------------------
    ws = wb.create_sheet("fields")
    header(ws, ["doc_id", "image", "field", "ground truth", "verdict", "correct value",
                "note"])
    n = 0
    for doc in chosen:
        rec = records[doc]
        for path in sorted(rec.annotated_fields):
            if path in ("lineItems", "totals.otherCharges"):
                continue
            v = rec.gt.get(path)
            ws.append([doc, f"images/{doc}.png", path,
                       "«absent»" if isinstance(v, _Absent) else
                       (json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list))
                        else str(v)),
                       "", "", ""])
            n += 1
    for r in range(2, n + 2):
        for c in (1, 2, 4):
            ws.cell(row=r, column=c).font = MONO
        ws.cell(row=r, column=3).font = BODY
    finish(ws, [30, 34, 30, 42, 14, 30, 40], [5, 6, 7], n)

    # ---- line_items ---------------------------------------------------------
    ws = wb.create_sheet("line_items")
    header(ws, ["doc_id", "image", "row #", "description", "itemCode", "qty", "unit price",
                "line total", "service date", "verdict", "correct value(s)",
                "rows missing from GT", "note"])
    n = 0
    for doc in chosen:
        rec = records[doc]
        rows = rec.gt.get("lineItems") or []
        if not rows:
            ws.append([doc, f"images/{doc}.png", "—", "«no itemised table in GT»",
                       "", "", "", "", "", "", "", "", ""])
            n += 1
            continue
        for i, r in enumerate(rows):
            ws.append([doc, f"images/{doc}.png", i,
                       r.get("description") or "", r.get("itemCode") or "",
                       r.get("quantity") or "", r.get("unitPrice") or "",
                       r.get("lineTotal") or "", r.get("serviceDate") or "",
                       "", "", "" if i else 0, ""])
            n += 1
    for r in range(2, n + 2):
        for c in (1, 2, 5, 6, 7, 8, 9):
            ws.cell(row=r, column=c).font = MONO
        ws.cell(row=r, column=4).font = BODY
    finish(ws, [30, 34, 7, 40, 24, 10, 12, 12, 14, 14, 30, 20, 34], [10, 11, 12, 13], n)

    wb.save(out)
    band_counts = collections.Counter(_band(int(records[d].meta.get("n_gt_rows") or 0))
                                      for d in chosen)
    total_rows = sum(int(records[d].meta.get("n_gt_rows") or 0) for d in chosen)
    print(f"written : {out}")
    print(f"sample  : {len(chosen)} documents, {total_rows} line-item rows")
    print(f"strata  : {dict(band_counts)}")
    print(f"images  : {adapter.root / 'images'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
